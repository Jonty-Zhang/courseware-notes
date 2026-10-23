#!/usr/bin/env python3
"""Build a Chinese academic .docx from a notes Markdown file.

Pipeline: Markdown --(preprocess)--> pandoc (reference.docx with the styles below)
--> python-docx post-processing.

Markdown conventions (see references/markdown-conventions.md):
  # 一、...   ## （一）...   ### 1. ...   #### （1）...   headings carry their own numbers
  $...$ / $$...$$          LaTeX math -> editable Word equations
  ==text==                 key point, rendered light blue (not bold)
  ![caption](fig.png){width=55%}   figure; "图 N " is prefixed automatically
  Table: caption           pipe-table caption; "表 N " is prefixed automatically

Formatting follows the word-format spec: A4, 2.5 cm margins, 宋体五号 body with
Times New Roman for Latin text, 1.5 line spacing, 2-char first-line indent,
一级标题 宋体小三加粗, 二级 黑体小四, 三级及以下 黑体五号, 楷体五号 captions,
centered native tables with 0 spacing in cells, centered page numbers.

  ::: example ... :::      worked example; dropped by --drop-examples

Usage: build_docx.py ch1.md [ch2.md ...] -o out.docx [--toc] [--drop-examples]
                     [--resource-path DIR]
"""
import argparse
import re
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

warnings.filterwarnings("ignore", message="style lookup by style_id is deprecated")

LATIN = "Times New Roman"
SONG, HEI, KAI = "宋体", "黑体", "楷体"
BODY_PT = 10.5          # 五号
KEY_COLOR = RGBColor(0x1F, 0x7A, 0xC4)   # light blue for key points
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


# ---------------------------------------------------------------- preprocess
EXAMPLE_BLOCK = re.compile(r"^:::\s*\{?\.?example\}?\s*$.*?^:::\s*$\n?", re.M | re.S)


def drop_examples(md: str) -> str:
    """Remove ::: example ... ::: blocks (the method-only edition)."""
    return EXAMPLE_BLOCK.sub("", md)


def preprocess(md: str) -> str:
    fig = 0
    tab = 0

    def figure(m):
        nonlocal fig
        fig += 1
        return f"![图 {fig} {m.group(1).strip()}]({m.group(2)})"

    def table(m):
        nonlocal tab
        tab += 1
        return f"{m.group(1)}Table: 表 {tab} {m.group(2).strip()}"

    out_lines = []
    in_math = False
    for line in md.splitlines():
        s = line.strip()
        opens = s.startswith("$$") and not (len(s) > 2 and s.endswith("$$"))
        closes = s.endswith("$$") and not s.startswith("$$")
        if opens or closes:          # a display block spread over several lines
            in_math = not in_math
        if not in_math:
            line = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", figure, line)
            line = re.sub(r"^(\s*)Table:\s*(.+)$", table, line)
            line = re.sub(r"==([^=\n]+?)==", r'[\1]{custom-style="重点"}', line)
        out_lines.append(line)
    return "\n".join(out_lines) + "\n"


# ---------------------------------------------------------------- styles
def set_fonts(style_or_run, east, latin=LATIN, size=None, bold=None, color=None):
    font = style_or_run.font
    font.name = latin
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), latin)
    rfonts.set(qn("w:eastAsia"), east)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color


def para_format(style, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0,
                spacing=1.5, keep_next=False):
    pf = style.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    pf.first_line_indent = Pt(BODY_PT * 2) if indent else Pt(0)
    pf.left_indent = Pt(0)
    if keep_next:
        pf.keep_with_next = True
    ppr = style.element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is not None:
        if indent:
            ind.set(qn("w:firstLineChars"), "200")
        else:
            for a in ("w:firstLineChars", "w:firstLine", "w:hangingChars", "w:hanging"):
                if ind.get(qn(a)) is not None:
                    del ind.attrib[qn(a)]


def get_style(doc, name, kind=1):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, kind)


def make_reference(path: Path):
    data = subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"],
                          check=True)
    doc = Document(str(path))
    st = doc.styles

    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        try:
            s = st[name]
        except KeyError:
            continue
        set_fonts(s, SONG, size=BODY_PT, color=RGBColor(0, 0, 0))
        para_format(s, indent=name not in ("Compact",))

    heads = {
        "Heading 1": (SONG, 15, True, 12, 6),
        "Heading 2": (HEI, 12, False, 9, 3),
        "Heading 3": (HEI, BODY_PT, False, 6, 0),
        "Heading 4": (HEI, BODY_PT, False, 3, 0),
        "Heading 5": (HEI, BODY_PT, False, 3, 0),
    }
    for name, (east, size, bold, before, after) in heads.items():
        s = st[name]
        set_fonts(s, east, size=size, bold=bold, color=RGBColor(0, 0, 0))
        s.font.italic = False
        para_format(s, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=before, after=after,
                    keep_next=True)
    st["Heading 1"].paragraph_format.page_break_before = True   # each chapter starts a page

    for name in ("Title", "Subtitle", "TOC Heading"):
        try:
            s = st[name]
        except KeyError:
            continue
        set_fonts(s, SONG, size=15, bold=True, color=RGBColor(0, 0, 0))
        para_format(s, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6)

    for name in ("Image Caption", "Table Caption", "Caption"):
        s = get_style(doc, name)
        set_fonts(s, KAI, size=BODY_PT, bold=False, color=RGBColor(0, 0, 0))
        s.font.italic = False
        para_format(s, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, spacing=1.2)
    for name in ("Captioned Figure", "Figure"):
        s = get_style(doc, name)
        para_format(s, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0, spacing=1.0,
                    keep_next=True)
    st["Table Caption"].paragraph_format.keep_with_next = True

    for level in (1, 2, 3):
        s = get_style(doc, f"toc {level}")
        set_fonts(s, SONG if level > 1 else HEI, size=BODY_PT, bold=False, color=RGBColor(0, 0, 0))
        para_format(s, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.25)
        s.paragraph_format.left_indent = Pt(BODY_PT * 2 * (level - 1))

    key = get_style(doc, "重点", kind=2)
    key.font.color.rgb = KEY_COLOR
    key.font.bold = False

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Cm(2.5))
    sec.gutter = Cm(0)
    doc.save(str(path))


# ---------------------------------------------------------------- post-process
def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_fonts(run, SONG, size=9)


def only_math(p):
    has_math = p._p.find(f"{{{M_NS}}}oMathPara") is not None
    text = "".join(t.text or "" for t in p._p.iter(qn("w:t")) if t.getparent().tag != f"{{{M_NS}}}r")
    return has_math and not text.strip()


def format_tables(doc):
    usable = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for t in doc.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        for old in tblpr.findall(qn("w:tblW")):
            tblpr.remove(old)
        w = OxmlElement("w:tblW")
        w.set(qn("w:w"), str(int(usable.twips if hasattr(usable, "twips") else usable / 635)))
        w.set(qn("w:type"), "dxa")
        tblpr.append(w)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "000000")
            borders.append(e)
        for old in tblpr.findall(qn("w:tblBorders")):
            tblpr.remove(old)
        tblpr.append(borders)
        for old in tblpr.findall(qn("w:tblLook")):
            tblpr.remove(old)
        for ri, row in enumerate(t.rows):
            if ri == 0:
                trpr = row._tr.get_or_add_trPr()
                hdr = OxmlElement("w:tblHeader")
                trpr.append(hdr)
            for cell in row.cells:
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    pf.first_line_indent = Pt(0)
                    pf.line_spacing = 1.15
                    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        set_fonts(r, SONG, size=BODY_PT, bold=(ri == 0) or None)


def fill_toc(doc):
    """Pre-fill the TOC field's cached result with the heading list.

    pandoc leaves the field empty, so Word/WPS would show nothing (or an English
    hint) until fields are updated. The field stays marked dirty and
    updateFields is set, so Word refreshes it with page numbers on open.
    """
    body = doc.element.body
    instr = next((i for i in body.iter(qn("w:instrText")) if (i.text or "").strip().startswith("TOC")), None)
    if instr is None:
        return
    run = instr.getparent()
    p0 = run.getparent()
    end = [c for c in run if c.tag == qn("w:fldChar") and c.get(qn("w:fldCharType")) == "end"]
    for e in end:
        run.remove(e)
    levels = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3}
    heads = [(levels[p.style.name], p.text.strip()) for p in doc.paragraphs
             if p.style.name in levels and p.text.strip()]
    styles = {lvl: doc.styles[f"toc {lvl}"].style_id for lvl in (1, 2, 3)}

    def set_style(p, lvl):
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            p.insert(0, ppr)
        for old in ppr.findall(qn("w:pStyle")):
            ppr.remove(old)
        ps = OxmlElement("w:pStyle")
        ps.set(qn("w:val"), styles[lvl])
        ppr.insert(0, ps)

    def text_run(text):
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        return r

    last = p0
    for i, (lvl, text) in enumerate(heads):
        if i == 0:
            p = p0
        else:
            p = OxmlElement("w:p")
            last.addnext(p)
        set_style(p, lvl)
        p.append(text_run(text))
        last = p
    r_end = OxmlElement("w:r")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r_end.append(fc)
    last.append(r_end)

    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)


def postprocess(path: Path):
    doc = Document(str(path))
    fill_toc(doc)
    for p in doc.paragraphs:
        if only_math(p):
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p._p.find(".//" + qn("w:drawing")) is not None:
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
    format_tables(doc)
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, side, Cm(2.5))
        add_page_number(sec)
    doc.save(str(path))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("markdown", nargs="+", help="one or more Markdown files, joined in order")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--toc", action="store_true", help="insert an automatic Word TOC (levels 1-3)")
    ap.add_argument("--drop-examples", action="store_true", help="method-only edition")
    ap.add_argument("--resource-path", default=None, help="folder that figure paths are relative to")
    args = ap.parse_args()

    srcs = [Path(p).resolve() for p in args.markdown]
    out = Path(args.output).resolve()
    res = Path(args.resource_path).resolve() if args.resource_path else srcs[0].parent
    md = "\n\n".join(p.read_text(encoding="utf-8") for p in srcs)
    if args.drop_examples:
        md = drop_examples(md)
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        ref = td / "reference.docx"
        make_reference(ref)
        pre = td / "pre.md"
        pre.write_text(preprocess(md), encoding="utf-8")
        cmd = ["pandoc", str(pre), "-f", "markdown+tex_math_dollars+bracketed_spans+fenced_divs",
               "-o", str(out), "--reference-doc", str(ref), "--resource-path", str(res)]
        if args.toc:
            cmd += ["--toc", "--toc-depth=3", "-M", "toc-title=目录"]
        subprocess.run(cmd, check=True)
    postprocess(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    sys.exit(main())
