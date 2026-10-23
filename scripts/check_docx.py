#!/usr/bin/env python3
"""Structural quality gate for a generated notes .docx (no renderer needed).

Checks against the word-format spec:
  page A4 and 2.5 cm margins; style fonts/sizes for body and headings;
  0 spacing in table cells; equations present as OMML (not images, not raw $);
  key-point (light blue) text under 20 %; no process/AI wording; figures have
  captions; heading numbering looks continuous.

    check_docx.py notes.docx [--source notes.md ...]

With --source, also compares the number of math spans in the Markdown with the
number of Word equations, which catches formulas pandoc failed to convert.
Exit status 1 if any check fails.
"""
import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
BANNED = ["根据用户要求", "AI 生成", "AI生成", "以下是整理", "这里插入", "作为AI", "如下所示：", "TODO"]


def style_font(style):
    rpr = style.element.rPr
    east = rpr.rFonts.get(qn("w:eastAsia")) if rpr is not None and rpr.rFonts is not None else None
    size = style.font.size.pt if style.font.size else None
    return east, size, style.font.bold


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--source", nargs="*", default=[])
    ap.add_argument("--drop-examples", action="store_true",
                    help="source was built with --drop-examples")
    args = ap.parse_args()

    doc = Document(args.docx)
    fails, notes = [], []

    sec = doc.sections[0]
    if abs(sec.page_width - Cm(21)) > Cm(0.1) or abs(sec.page_height - Cm(29.7)) > Cm(0.1):
        fails.append("page is not A4")
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        if abs(getattr(sec, side) - Cm(2.5)) > Cm(0.05):
            fails.append(f"{side} is not 2.5 cm")

    expect = {"Body Text": ("宋体", 10.5, None), "Heading 1": ("宋体", 15, True),
              "Heading 2": ("黑体", 12, False), "Heading 3": ("黑体", 10.5, False)}
    for name, (east, size, bold) in expect.items():
        try:
            got = style_font(doc.styles[name])
        except KeyError:
            continue
        if got[0] != east or got[1] != size or (bold is not None and bool(got[2]) != bold):
            fails.append(f"style {name}: got {got}, want {(east, size, bold)}")

    for ti, t in enumerate(doc.tables):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    if (pf.space_after or 0) != 0 or (pf.space_before or 0) != 0:
                        fails.append(f"table {ti + 1}: cell paragraph has spacing")
                        break

    body = doc.element.body
    n_math = len(body.findall(f".//{M}oMath"))
    text = "".join(t.text or "" for t in body.iter(qn("w:t")))
    if re.search(r"(?<!\\)\$[^$\s][^$]*\$", text):
        fails.append("raw $...$ left in text (math not converted)")
    for w in BANNED:
        if w in text:
            fails.append(f"process wording found: {w}")

    key_chars = 0
    for r in body.iter(qn("w:r")):
        rs = r.find(qn("w:rPr") + "/" + qn("w:rStyle"))
        if rs is not None and rs.get(qn("w:val")) in ("重点", "a7", "Key"):
            key_chars += sum(len(t.text or "") for t in r.iter(qn("w:t")))
    ratio = key_chars / max(1, len(text))
    if ratio > 0.2:
        fails.append(f"key-point text is {ratio:.0%} (> 20%)")

    n_img = len(body.findall(".//" + qn("w:drawing")))
    caps = [p.text for p in doc.paragraphs if p.style.name in ("Image Caption", "Caption")]
    fig_caps = [c for c in caps if c.startswith("图")]
    if n_img and len(fig_caps) < n_img:
        fails.append(f"{n_img} images but only {len(fig_caps)} figure captions")

    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    cn = "一二三四五六七八九十"
    for i, h in enumerate(h1):
        if i < len(cn) and not h.startswith(cn[i] + "、"):
            notes.append(f"heading 1 #{i + 1} numbering looks off: {h[:20]}")

    if args.source:
        sys.path.insert(0, str(Path(__file__).parent))
        from build_docx import drop_examples
        md = "\n\n".join(Path(s).read_text(encoding="utf-8") for s in args.source)
        if args.drop_examples:
            md = drop_examples(md)
        display = len(re.findall(r"\$\$.+?\$\$", md, re.S))
        inline = len(re.findall(r"(?<!\$)\$(?!\$)[^$\n]+?(?<!\$)\$(?!\$)", re.sub(r"\$\$.+?\$\$", "", md, flags=re.S)))
        if n_math < display + inline:
            fails.append(f"Markdown has {display + inline} math spans, Word has {n_math} equations")
        notes.append(f"math spans in source: {display + inline}")

    print(f"equations: {n_math}  images: {n_img}  tables: {len(doc.tables)}  "
          f"key text: {ratio:.1%}  chars: {len(text)}")
    for n in notes:
        print("note:", n)
    for f in fails:
        print("FAIL:", f)
    print("PASS" if not fails else f"{len(fails)} problem(s)")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
